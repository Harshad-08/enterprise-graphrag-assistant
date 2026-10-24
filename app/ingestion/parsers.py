"""Format-specific text extraction for PDF (PyMuPDF), DOCX (python-docx) and TXT."""

from __future__ import annotations

import io
import logging
from pathlib import Path

from app.models.schemas import ParsedPage
from app.utils.text import normalize_whitespace

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


class DocumentParseError(ValueError):
    """Raised when a file cannot be parsed into text."""


def detect_format(filename: str, content: bytes) -> str:
    """Validate the extension *and* the file signature so a renamed binary is rejected."""
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise DocumentParseError(f"Unsupported file type '{ext}'. Allowed: PDF, DOCX, TXT.")
    if ext == ".pdf" and not content.startswith(b"%PDF"):
        raise DocumentParseError("File has a .pdf extension but is not a valid PDF.")
    if ext == ".docx" and not content.startswith(b"PK"):
        raise DocumentParseError("File has a .docx extension but is not a valid DOCX archive.")
    if ext == ".txt":
        if b"\x00" in content[:4096]:
            raise DocumentParseError("TXT file appears to be binary.")
    return ext.lstrip(".")


def parse_txt(content: bytes) -> list[ParsedPage]:
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - latin-1 always decodes
        raise DocumentParseError("Could not decode text file.")
    return [ParsedPage(text=normalize_whitespace(text), page_number=None)]


def parse_pdf(content: bytes) -> list[ParsedPage]:
    import pymupdf as fitz

    try:
        doc = fitz.open(stream=content, filetype="pdf")
    except Exception as exc:  # PyMuPDF raises several internal error types
        raise DocumentParseError(f"Could not open PDF: {exc}") from exc
    pages = []
    with doc:
        for index, page in enumerate(doc):
            text = normalize_whitespace(page.get_text("text"))
            if text:
                pages.append(ParsedPage(text=text, page_number=index + 1))
    if not pages:
        raise DocumentParseError("PDF contains no extractable text (scanned PDFs need OCR).")
    return pages


def parse_docx(content: bytes) -> list[ParsedPage]:
    import docx

    try:
        document = docx.Document(io.BytesIO(content))
    except Exception as exc:
        raise DocumentParseError(f"Could not open DOCX: {exc}") from exc
    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            lines.append("")
            continue
        style = (paragraph.style.name or "").lower() if paragraph.style is not None else ""
        # Preserve heading structure so the chunker can attach section metadata.
        if style.startswith("heading") and not text.startswith("#"):
            text = f"# {text}"
        lines.append(text)
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    text = normalize_whitespace("\n".join(lines))
    if not text:
        raise DocumentParseError("DOCX contains no text.")
    return [ParsedPage(text=text, page_number=None)]


def parse_document(filename: str, content: bytes) -> tuple[str, list[ParsedPage]]:
    fmt = detect_format(filename, content)
    parser = {"pdf": parse_pdf, "docx": parse_docx, "txt": parse_txt}[fmt]
    pages = parser(content)
    if not any(p.text.strip() for p in pages):
        raise DocumentParseError("Document is empty.")
    logger.debug("Parsed %s (%s) into %d page(s)", filename, fmt, len(pages))
    return fmt, pages
