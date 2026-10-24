"""Parsing, metadata extraction and chunking."""

from __future__ import annotations

import io
from datetime import date

import docx
import pymupdf
import pytest

from app.ingestion.chunking import chunk_document
from app.ingestion.metadata import build_metadata, split_front_matter, strip_front_matter
from app.ingestion.parsers import DocumentParseError, detect_format, parse_document
from app.models.schemas import DocumentStatus, ParsedPage

TXT = b"""---
document_id: TST-001
title: Test Policy
doc_type: policy
version: 1.0
publication_date: 2024-01-01
effective_date: 2024-02-01
status: active
authority: department
supersedes: TST-000
---
# 1. Scope
This policy governs the Production Database.

# 2. Rules
Passwords must contain at least 14 characters. Access must be approved.
"""


def _pdf_bytes(pages: list[str]) -> bytes:
    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        page.insert_textbox(pymupdf.Rect(50, 50, 550, 800), text, fontsize=10)
    data = doc.tobytes()
    doc.close()
    return data


def _docx_bytes() -> bytes:
    document = docx.Document()
    document.add_heading("Access Rules", level=1)
    document.add_paragraph("Engineers must use the Bastion Host.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class TestParsers:
    def test_txt_parsing(self):
        fmt, pages = parse_document("policy.txt", TXT)
        assert fmt == "txt"
        assert len(pages) == 1 and pages[0].page_number is None
        assert "Passwords must contain" in pages[0].text

    def test_pdf_parsing_keeps_page_numbers(self):
        fmt, pages = parse_document("doc.pdf", _pdf_bytes(["First page text.", "Second page text."]))
        assert fmt == "pdf"
        assert [p.page_number for p in pages] == [1, 2]
        assert "Second page" in pages[1].text

    def test_docx_parsing_marks_headings(self):
        fmt, pages = parse_document("doc.docx", _docx_bytes())
        assert fmt == "docx"
        assert "# Access Rules" in pages[0].text
        assert pages[0].page_number is None

    @pytest.mark.parametrize(
        "filename,content",
        [("x.exe", b"MZ"), ("fake.pdf", b"not a pdf"), ("fake.docx", b"plain text"), ("bin.txt", b"\x00\x01\x02")],
    )
    def test_rejects_invalid_files(self, filename, content):
        with pytest.raises(DocumentParseError):
            detect_format(filename, content)


class TestMetadata:
    def test_front_matter_is_parsed_and_removed(self):
        fields, body = split_front_matter(TXT.decode())
        assert fields["document_id"] == "TST-001"
        assert not body.lstrip().startswith("---")

    def test_build_metadata_from_front_matter(self):
        fields, pages = strip_front_matter([ParsedPage(text=TXT.decode())])
        meta = build_metadata(fields, filename="p.txt", source_format="txt", body_text=pages[0].text,
                              known_authorities={"department", "executive"})
        assert meta.document_id == "TST-001"
        assert meta.effective_date == date(2024, 2, 1)
        assert meta.status == DocumentStatus.ACTIVE
        assert meta.supersedes == ["TST-000"]
        assert meta.metadata_warnings == []

    def test_missing_and_invalid_metadata_is_flagged_not_guessed(self):
        meta = build_metadata(
            {"effective_date": "31/12/2024", "status": "current", "authority": "ceo-ish"},
            filename="Untitled Doc.txt", source_format="txt", body_text="# Heading\nSome text.",
            known_authorities={"executive"},
        )
        assert meta.effective_date is None
        assert meta.status == DocumentStatus.UNKNOWN
        assert meta.authority is None
        joined = " ".join(meta.metadata_warnings)
        assert "Invalid effective_date" in joined and "Unknown authority" in joined and "Missing publication_date" in joined

    def test_generated_document_id_is_stable(self):
        a = build_metadata({}, filename="a.txt", source_format="txt", body_text="same body")
        b = build_metadata({}, filename="a.txt", source_format="txt", body_text="same body")
        c = build_metadata({}, filename="a.txt", source_format="txt", body_text="different body")
        assert a.document_id == b.document_id != c.document_id

    def test_overrides_win_and_self_supersession_is_removed(self):
        meta = build_metadata({"document_id": "X-1", "title": "Old"}, filename="x.txt", source_format="txt",
                              body_text="t", overrides={"title": "New", "supersedes": "X-1;X-0"})
        assert meta.title == "New"
        assert meta.supersedes == ["X-0"]


class TestChunking:
    def test_chunks_have_stable_ids_sections_and_pages(self):
        pages = [ParsedPage(text="# 1. Scope\nAlpha sentence one. Alpha two.\n\n# 2. Rules\nBeta rule.", page_number=3)]
        chunks = chunk_document("DOC-1", pages, max_chars=500)
        assert [c.chunk_id for c in chunks] == ["DOC-1::c000", "DOC-1::c001"]
        assert [c.section for c in chunks] == ["1. Scope", "2. Rules"]
        assert all(c.page == 3 for c in chunks)
        assert chunk_document("DOC-1", pages, max_chars=500) == chunks  # deterministic

    def test_long_sections_are_split_with_overlap(self):
        text = " ".join(f"Sentence number {i} is about access control rules." for i in range(30))
        chunks = chunk_document("D", [ParsedPage(text=text)], max_chars=200, overlap_sentences=1)
        assert len(chunks) > 3
        assert all(len(c.text) <= 260 for c in chunks)
        last_sentence_of_first = chunks[0].text.split(". ")[-1]
        assert last_sentence_of_first.rstrip(".") in chunks[1].text

    def test_chunks_never_span_pages(self):
        pages = [ParsedPage(text="Page one text.", page_number=1), ParsedPage(text="Page two text.", page_number=2)]
        chunks = chunk_document("D", pages)
        assert [c.page for c in chunks] == [1, 2]
