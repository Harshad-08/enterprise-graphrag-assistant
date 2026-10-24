"""Render the synthetic sources in data/synthetic_sources/ as PDF and DOCX.

TECH-DB-002 becomes a two-page PDF (exercises PyMuPDF and page numbers) and
PROC-PART-003 becomes a DOCX with heading styles (exercises python-docx).
The outputs are written to data/synthetic_documents/.
"""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

SOURCES = ROOT / "data" / "synthetic_sources"
OUTPUT = ROOT / "data" / "synthetic_documents"
PAGE_BREAK = "=== PAGE BREAK ==="


def write_pdf(source: Path, target: Path) -> None:
    import pymupdf as fitz

    doc = fitz.open()
    for page_text in source.read_text().split(PAGE_BREAK):
        page = doc.new_page(width=595, height=842)  # A4
        rect = fitz.Rect(56, 56, 539, 786)
        overflow = page.insert_textbox(rect, page_text.strip(), fontsize=10, fontname="helv")
        if overflow < 0:
            raise RuntimeError(f"Text overflowed the page in {source.name}")
    doc.save(target)
    doc.close()


def write_docx(source: Path, target: Path) -> None:
    import docx

    document = docx.Document()
    for line in source.read_text().splitlines():
        if line.startswith("# "):
            document.add_heading(line[2:].strip(), level=1)
        else:
            document.add_paragraph(line)
    document.save(target)


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    write_pdf(SOURCES / "TECH-DB-002_database_platform_overview.txt", OUTPUT / "TECH-DB-002_database_platform_overview.pdf")
    write_docx(SOURCES / "PROC-PART-003_partner_integration_onboarding.txt", OUTPUT / "PROC-PART-003_partner_integration_onboarding.docx")
    print("Generated PDF and DOCX synthetic documents in", OUTPUT)


if __name__ == "__main__":
    main()
